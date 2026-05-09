#!/usr/bin/env python3
"""
audit_april2026_language_mixups.py — audit April-2026 FR/Adja column mixups.

This script is intentionally audit-only. It reproduces the current
ingest_april2026.py parsing in memory, then writes a review TSV of rows that
look like French/Adja boundary mistakes, swaps, missing targets, or unresolved
translator placeholders. It does not rewrite new6k.tsv or any split files.

Run from repo root:
    python experiments/data/audit_april2026_language_mixups.py
"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pandas as pd

import ingest_april2026 as ingest


REPO_ROOT = Path(__file__).resolve().parent.parent.parent
DEFAULT_RAW = REPO_ROOT / "april-2026-new-data" / "raw"
DEFAULT_OUT = REPO_ROOT / "april-2026-new-data" / "processed"
DEFAULT_AUDIT = DEFAULT_OUT / "mixed_language_audit.tsv"
DEFAULT_NEEDS_REVIEW = DEFAULT_OUT / "needs_review.tsv"

# Keep this aligned with ingest_april2026.py. These are strong evidence that
# the French column contains Adja text.
ADJA_CHARS = set("ɖɛɔŋɣʋɥʔᴐ")

# Broad token matcher for French/Adja Latin text. It includes the Adja-specific
# characters above plus Latin extended letters that appear in these batches.
TOKEN_RE = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿƷƸɖɛɔŋɣʋɥʔᴐ'’.-]+")

FRENCH_DIACRITICS = set("àâçéèêëîïôùûüÿœæ")

# Function words and content words that should normally stay on the French side
# when they appear as an Adja-column prefix before the real Adja translation.
FRENCH_CONTINUATION_WORDS = {
    "a",
    "à",
    "au",
    "aux",
    "avec",
    "ce",
    "ces",
    "cet",
    "cette",
    "comme",
    "contre",
    "d",
    "dans",
    "de",
    "des",
    "du",
    "en",
    "et",
    "la",
    "le",
    "les",
    "lui",
    "ma",
    "me",
    "mes",
    "mon",
    "nous",
    "ou",
    "où",
    "pas",
    "pour",
    "qu",
    "que",
    "sa",
    "sans",
    "se",
    "ses",
    "son",
    "sur",
    "te",
    "ton",
    "une",
    "un",
}

ADJA_SHORT_STARTERS = {
    "a",
    "be",
    "da",
    "de",
    "e",
    "ji",
    "le",
    "mi",
    "na",
    "nɔ",
    "sa",
    "se",
    "va",
    "wo",
    "yi",
}

INCOMPLETE_FRENCH_ENDINGS = {
    "a",
    "à",
    "au",
    "aux",
    "avec",
    "ce",
    "ces",
    "cet",
    "cette",
    "comme",
    "contre",
    "dans",
    "de",
    "des",
    "du",
    "en",
    "et",
    "la",
    "le",
    "les",
    "ma",
    "mes",
    "mon",
    "notre",
    "ou",
    "où",
    "pour",
    "qu",
    "que",
    "sa",
    "sans",
    "son",
    "sur",
    "ton",
    "une",
    "un",
}

PLACEHOLDER_RE = re.compile(r"\?{2,}|…|\.{3,}")


@dataclass
class ParsedRow:
    source: str
    raw_line: str
    current_french: str
    current_adja: str
    parse_method: str
    parse_confidence: float | None = None


@dataclass
class AuditIssue:
    source: str
    raw_line: str
    current_french: str
    current_adja: str
    issue_type: str
    suggested_french: str
    suggested_adja: str
    confidence_bucket: str
    detail: str
    parse_method: str
    parse_confidence: float | None


def normal_token(token: str) -> str:
    return token.lower().replace("’", "'").strip(".,;:!?()[]{}\"“”")


def token_spans(text: str) -> list[tuple[str, int, int]]:
    out = []
    for m in TOKEN_RE.finditer(text):
        token = normal_token(m.group(0))
        if token:
            out.append((token, m.start(), m.end()))
    return out


def has_adja_char(text: str) -> bool:
    return any(ch in ADJA_CHARS for ch in text.lower())


def count_adja_chars(text: str) -> int:
    return sum(1 for ch in text.lower() if ch in ADJA_CHARS)


def has_french_diacritic(token: str) -> bool:
    return any(ch in FRENCH_DIACRITICS for ch in token.lower())


def has_placeholder(text: str) -> bool:
    return bool(PLACEHOLDER_RE.search(text))


def french_ending_is_incomplete(text: str) -> str | None:
    spans = token_spans(text)
    if not spans:
        return None
    last = spans[-1][0]
    if last in INCOMPLETE_FRENCH_ENDINGS:
        return last
    return None


def token_looks_french_continuation(
    token: str,
    fr_vocab: set[str],
    aj_vocab: set[str],
) -> bool:
    if not token:
        return False
    if has_adja_char(token):
        return False
    if token in FRENCH_CONTINUATION_WORDS:
        return True
    if has_french_diacritic(token):
        return True
    if "'" in token and token.split("'", 1)[0] in {"c", "d", "j", "l", "m", "n", "qu", "s", "t"}:
        return True
    if token in fr_vocab and token not in aj_vocab:
        return True
    # Conservative fallback for long French-looking words before the first
    # Adja-specific character. This catches examples such as "contribution",
    # "croustillant", and "villages" without relying on an external language ID.
    return len(token) >= 7 and token not in aj_vocab


def leading_french_prefix(
    text: str,
    fr_vocab: set[str],
    aj_vocab: set[str],
) -> tuple[str, int, list[str]]:
    """Return likely French continuation at the start of an Adja field.

    The returned int is the character offset after the prefix, suitable for
    moving that text back to the French side.
    """
    spans = token_spans(text)
    prefix_tokens: list[str] = []
    prefix_end = 0
    allow_content_after_connector = False

    for token, _start, end in spans:
        if has_adja_char(token):
            break

        if token in FRENCH_CONTINUATION_WORDS:
            prefix_tokens.append(token)
            prefix_end = end
            allow_content_after_connector = True
            continue

        if token_looks_french_continuation(token, fr_vocab, aj_vocab) and (
            not prefix_tokens
            or allow_content_after_connector
            or has_french_diacritic(token)
            or ("'" in token and token.split("'", 1)[0] in {"c", "d", "j", "l", "m", "n", "qu", "s", "t"})
            or (token in fr_vocab and token not in aj_vocab)
        ):
            prefix_tokens.append(token)
            prefix_end = end
            allow_content_after_connector = False
            continue

        # After a strong French prefix starts, allow short French connectors to
        # stay attached. This handles phrases such as "la crotte de cochon".
        if prefix_tokens and token in FRENCH_CONTINUATION_WORDS:
            prefix_tokens.append(token)
            prefix_end = end
            allow_content_after_connector = True
            continue

        # Also allow one or more French-looking content words after a determiner
        # or preposition. This catches "le champ" in "sur le champ" without
        # pulling in obvious short Adja starters such as "e", "wo", or "mi".
        if prefix_tokens and allow_content_after_connector and len(token) >= 4 and token not in ADJA_SHORT_STARTERS:
            prefix_tokens.append(token)
            prefix_end = end
            allow_content_after_connector = False
            continue

        break

    if not prefix_tokens:
        return "", 0, []
    return text[:prefix_end].strip(), prefix_end, prefix_tokens


def adja_evidence(text: str, aj_vocab: set[str]) -> int:
    hits = count_adja_chars(text)
    hits += sum(1 for token, _start, _end in token_spans(text) if token in aj_vocab)
    return hits


def suggest_move_prefix(row: ParsedRow, prefix_end: int) -> tuple[str, str]:
    while prefix_end < len(row.current_adja) and row.current_adja[prefix_end] in ")]}":
        prefix_end += 1
    moved = row.current_adja[:prefix_end].strip()
    remaining = row.current_adja[prefix_end:].strip(" \t-—–|")
    suggested_fr = f"{row.current_french} {moved}".strip()
    return suggested_fr, remaining


def choose_prefix_confidence(
    row: ParsedRow,
    prefix_tokens: list[str],
    suggested_adja: str,
    incomplete_ending: str | None,
    aj_vocab: set[str],
) -> str:
    if not suggested_adja:
        return "needs-human"
    evidence = adja_evidence(suggested_adja, aj_vocab)
    strong_prefix = len(prefix_tokens) >= 2 or any(has_french_diacritic(t) for t in prefix_tokens)
    if evidence > 0 and (strong_prefix or incomplete_ending is not None):
        return "high"
    if evidence > 0:
        return "medium"
    return "needs-human"


def audit_row(row: ParsedRow, fr_vocab: set[str], aj_vocab: set[str]) -> list[AuditIssue]:
    issues: list[AuditIssue] = []

    fr_adja_chars = count_adja_chars(row.current_french)
    if fr_adja_chars:
        issues.append(
            AuditIssue(
                source=row.source,
                raw_line=row.raw_line,
                current_french=row.current_french,
                current_adja=row.current_adja,
                issue_type="adja_chars_in_french_column",
                suggested_french="",
                suggested_adja="",
                confidence_bucket="high",
                detail=f"{fr_adja_chars} Adja-specific character(s) found in current French field",
                parse_method=row.parse_method,
                parse_confidence=row.parse_confidence,
            )
        )

    # Boundary heuristics apply only to the no-separator HUREINE-style parser.
    # Colon/table/tabular sources already provide an explicit FR/AJ boundary, so
    # prefix-like Adja text such as "le ..." should not be treated as a split bug.
    if row.parse_method == "lexicon_split":
        incomplete = french_ending_is_incomplete(row.current_french)
        prefix, prefix_end, prefix_tokens = leading_french_prefix(row.current_adja, fr_vocab, aj_vocab)
        if prefix:
            suggested_fr, suggested_aj = suggest_move_prefix(row, prefix_end)
            issues.append(
                AuditIssue(
                    source=row.source,
                    raw_line=row.raw_line,
                    current_french=row.current_french,
                    current_adja=row.current_adja,
                    issue_type="french_prefix_in_adja_column",
                    suggested_french=suggested_fr,
                    suggested_adja=suggested_aj,
                    confidence_bucket=choose_prefix_confidence(
                        row,
                        prefix_tokens,
                        suggested_aj,
                        incomplete,
                        aj_vocab,
                    ),
                    detail=f"Adja column starts with likely French continuation: {prefix!r}",
                    parse_method=row.parse_method,
                    parse_confidence=row.parse_confidence,
                )
            )
        elif incomplete:
            issues.append(
                AuditIssue(
                    source=row.source,
                    raw_line=row.raw_line,
                    current_french=row.current_french,
                    current_adja=row.current_adja,
                    issue_type="incomplete_french_boundary",
                    suggested_french="",
                    suggested_adja="",
                    confidence_bucket="medium",
                    detail=f"Current French field ends with incomplete token {incomplete!r}",
                    parse_method=row.parse_method,
                    parse_confidence=row.parse_confidence,
                )
            )

    if has_placeholder(row.raw_line) or has_placeholder(row.current_french) or has_placeholder(row.current_adja):
        issues.append(
            AuditIssue(
                source=row.source,
                raw_line=row.raw_line,
                current_french=row.current_french,
                current_adja=row.current_adja,
                issue_type="placeholder_or_unknown_translation",
                suggested_french="",
                suggested_adja="",
                confidence_bucket="needs-human",
                detail="Question-mark/ellipsis placeholder detected",
                parse_method=row.parse_method,
                parse_confidence=row.parse_confidence,
            )
        )

    if not row.current_adja.strip():
        issues.append(
            AuditIssue(
                source=row.source,
                raw_line=row.raw_line,
                current_french=row.current_french,
                current_adja=row.current_adja,
                issue_type="empty_adja_after_parse",
                suggested_french="",
                suggested_adja="",
                confidence_bucket="needs-human",
                detail="Parser produced an empty Adja field",
                parse_method=row.parse_method,
                parse_confidence=row.parse_confidence,
            )
        )
    elif row.parse_method == "lexicon_split" and adja_evidence(row.current_adja, aj_vocab) == 0:
        issues.append(
            AuditIssue(
                source=row.source,
                raw_line=row.raw_line,
                current_french=row.current_french,
                current_adja=row.current_adja,
                issue_type="no_reliable_adja_evidence",
                suggested_french="",
                suggested_adja="",
                confidence_bucket="needs-human",
                detail="Lexicon split target has no Adja-specific characters or known Adja vocabulary hits",
                parse_method=row.parse_method,
                parse_confidence=row.parse_confidence,
            )
        )

    return issues


def parse_docx_rows(
    path: Path,
    source_label: str,
    fr_vocab: set[str],
    aj_vocab: set[str],
) -> Iterable[ParsedRow]:
    try:
        import docx
    except ImportError as e:
        raise RuntimeError("python-docx is required for DOCX audit") from e

    doc = docx.Document(str(path))

    table_rows: list[ParsedRow] = []
    for tbl in doc.tables:
        for r in tbl.rows:
            cells = [c.text for c in r.cells]
            if len(cells) < 2:
                continue
            fr, aj = cells[0], cells[1]
            if fr.strip().lower() in ingest.FR_COL_ALIASES and aj.strip().lower() in ingest.AJ_COL_ALIASES:
                continue
            table_rows.append(
                ParsedRow(
                    source=source_label,
                    raw_line="\t".join(ingest.clean(c) for c in cells[:2]),
                    current_french=ingest.clean(fr),
                    current_adja=ingest.clean(aj),
                    parse_method="table",
                    parse_confidence=None,
                )
            )

    # Mirror ingest.extract_docx(): if tables are present, paragraph parsing is
    # skipped entirely.
    if table_rows:
        yield from table_rows
        return

    for para in doc.paragraphs:
        raw = para.text.strip()
        if not raw:
            continue
        line = ingest._strip_leading_number(raw)
        if not line:
            continue

        if ":" in line:
            left, _, right = line.rpartition(":")
            left = left.strip(" \t-—–|")
            right = right.strip()
            if left and right:
                yield ParsedRow(
                    source=source_label,
                    raw_line=ingest.clean(line),
                    current_french=ingest.clean(left),
                    current_adja=ingest.clean(right),
                    parse_method="colon_split",
                    parse_confidence=None,
                )
            continue

        fr, aj, conf = ingest._lexicon_split(line, fr_vocab, aj_vocab)
        if fr and aj and conf >= 2.0:
            yield ParsedRow(
                source=source_label,
                raw_line=ingest.clean(line),
                current_french=ingest.clean(fr),
                current_adja=ingest.clean(aj),
                parse_method="lexicon_split",
                parse_confidence=conf,
            )


def parse_tabular_rows(path: Path, source_label: str) -> Iterable[ParsedRow]:
    df = ingest.extract_tabular(path)
    for _, row in df.iterrows():
        fr = ingest.clean(row["french"])
        aj = ingest.clean(row["adja"])
        yield ParsedRow(
            source=source_label,
            raw_line=f"{fr}\t{aj}",
            current_french=fr,
            current_adja=aj,
            parse_method="tabular",
            parse_confidence=None,
        )


def iter_current_parsed_rows(
    raw_dir: Path,
    fr_vocab: set[str],
    aj_vocab: set[str],
) -> Iterable[ParsedRow]:
    files = sorted(p for p in raw_dir.iterdir() if p.is_file() and not p.name.startswith("."))
    for path in files:
        ext = path.suffix.lower()
        source_label = path.name
        target = path

        if ext == ".doc":
            target = ingest._convert_doc_to_docx(path)
            ext = ".docx"

        if ext == ".docx":
            yield from parse_docx_rows(target, source_label, fr_vocab, aj_vocab)
        elif ext in {".csv", ".tsv", ".xlsx", ".xls"}:
            yield from parse_tabular_rows(target, source_label)
        else:
            print(f"skip unsupported file: {path.name}", file=sys.stderr)


def read_existing_needs_review(path: Path) -> list[AuditIssue]:
    if not path.exists():
        return []

    issues: list[AuditIssue] = []
    with path.open(encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f, delimiter="\t")
        for row in reader:
            issues.append(
                AuditIssue(
                    source=row.get("source", ""),
                    raw_line=row.get("line", ""),
                    current_french=row.get("best_french", ""),
                    current_adja=row.get("best_adja", ""),
                    issue_type="existing_low_confidence_split",
                    suggested_french="",
                    suggested_adja="",
                    confidence_bucket="needs-human",
                    detail="Previously emitted by ingest_april2026.py needs_review.tsv",
                    parse_method="low_confidence_lexicon_split",
                    parse_confidence=float(row["confidence"]) if row.get("confidence") else None,
                )
            )
    return issues


def dedupe_issues(issues: Iterable[AuditIssue]) -> list[AuditIssue]:
    seen: set[tuple[str, str, str, str, str]] = set()
    out: list[AuditIssue] = []
    for issue in issues:
        key = (
            issue.source,
            issue.raw_line,
            issue.current_french,
            issue.current_adja,
            issue.issue_type,
        )
        if key in seen:
            continue
        seen.add(key)
        out.append(issue)
    return out


def write_audit(issues: list[AuditIssue], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "source",
                "raw_line",
                "current_french",
                "current_adja",
                "issue_type",
                "suggested_french",
                "suggested_adja",
                "confidence_bucket",
                "detail",
                "parse_method",
                "parse_confidence",
            ],
            delimiter="\t",
        )
        writer.writeheader()
        for issue in issues:
            writer.writerow(
                {
                    "source": issue.source,
                    "raw_line": issue.raw_line,
                    "current_french": issue.current_french,
                    "current_adja": issue.current_adja,
                    "issue_type": issue.issue_type,
                    "suggested_french": issue.suggested_french,
                    "suggested_adja": issue.suggested_adja,
                    "confidence_bucket": issue.confidence_bucket,
                    "detail": issue.detail,
                    "parse_method": issue.parse_method,
                    "parse_confidence": "" if issue.parse_confidence is None else issue.parse_confidence,
                }
            )


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n", 1)[0])
    ap.add_argument("--raw-dir", type=Path, default=DEFAULT_RAW)
    ap.add_argument("--output", type=Path, default=DEFAULT_AUDIT)
    ap.add_argument("--needs-review", type=Path, default=DEFAULT_NEEDS_REVIEW)
    args = ap.parse_args()

    if not args.raw_dir.exists() or not any(args.raw_dir.iterdir()):
        print(f"ERROR: raw directory is empty or missing: {args.raw_dir}", file=sys.stderr)
        return 1

    if shutil.which("textutil") is None and any(p.suffix.lower() == ".doc" for p in args.raw_dir.iterdir()):
        print("ERROR: .doc audit requires macOS textutil for conversion", file=sys.stderr)
        return 1

    fr_vocab, aj_vocab = ingest.build_side_vocabs()
    print(f"Built side vocabs: |FR|={len(fr_vocab):,} |AJ|={len(aj_vocab):,}")

    parsed_count = 0
    issues: list[AuditIssue] = []
    for row in iter_current_parsed_rows(args.raw_dir, fr_vocab, aj_vocab):
        parsed_count += 1
        issues.extend(audit_row(row, fr_vocab, aj_vocab))

    existing = read_existing_needs_review(args.needs_review)
    issues.extend(existing)
    issues = dedupe_issues(issues)

    # Put the most actionable rows first for manual review.
    confidence_rank = {"high": 0, "medium": 1, "needs-human": 2}
    issues.sort(
        key=lambda x: (
            confidence_rank.get(x.confidence_bucket, 9),
            x.source,
            x.issue_type,
            x.raw_line,
        )
    )

    write_audit(issues, args.output)

    by_conf = Counter(i.confidence_bucket for i in issues)
    by_type = Counter(i.issue_type for i in issues)
    print(f"Parsed rows audited: {parsed_count:,}")
    print(f"Existing needs_review rows included: {len(existing):,}")
    print(f"Wrote audit rows: {len(issues):,} -> {args.output}")
    print("By confidence:")
    for key in ("high", "medium", "needs-human"):
        print(f"  {key}: {by_conf.get(key, 0):,}")
    print("By issue type:")
    for key, value in by_type.most_common():
        print(f"  {key}: {value:,}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
